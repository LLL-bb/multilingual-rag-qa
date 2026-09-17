"""多格式文档解析。

支持 TXT / MD / PDF / DOCX,统一成"段落块列表"这一种中间表示,
后续的分块、语言检测、嵌入都不需要关心原始格式。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED_SUFFIXES = {".txt", ".md", ".markdown", ".pdf", ".docx"}

# 顺序有讲究:utf-8-sig 必须排在 utf-8 前面。
# 带 BOM 的文件用 utf-8 也能解码成功,只是会把 ﻿ 留在正文开头 ——
# 一旦排在后面就永远轮不到它,污染第一个块的内容和嵌入。
# utf-8-sig 对不带 BOM 的 UTF-8 同样正确,所以放第一个没有代价。
_TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030")


@dataclass
class Document:
    """一篇文档解析后的结果。

    blocks 是段落级的文本块,是分块的输入单位。
    """

    doc_id: str
    source: str  # 相对于语料根目录的路径,用于在答案里引用出处
    fmt: str
    blocks: list[str] = field(default_factory=list)

    @property
    def char_count(self) -> int:
        return sum(len(b) for b in self.blocks)


def _split_blocks(text: str) -> list[str]:
    """按空行切段,去掉首尾空白和空段。"""
    return [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]


def _read_text(path: Path) -> str:
    for encoding in _TEXT_ENCODINGS:
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def _load_text(path: Path) -> list[str]:
    return _split_blocks(_read_text(path))


def _load_pdf(path: Path) -> list[str]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    blocks: list[str] = []
    for page in reader.pages:
        # PDF 抽取出来的文本经常丢段落边界,只能按空行尽力还原;
        # 还原不了也没关系,分块阶段还会按句子再切一次。
        blocks.extend(_split_blocks(page.extract_text() or ""))
    return blocks


def _load_docx(path: Path) -> list[str]:
    import docx

    document = docx.Document(str(path))
    blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    # 表格常承载关键信息(如课程清单、餐品价目),按行拼成一行文本保留下来。
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return blocks


_LOADERS = {
    ".txt": _load_text,
    ".md": _load_text,
    ".markdown": _load_text,
    ".pdf": _load_pdf,
    ".docx": _load_docx,
}


def load_document(path: Path, root: Path | None = None) -> Document:
    """解析单个文件。"""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in _LOADERS:
        raise ValueError(
            f"不支持的格式: {path.name}(支持 {sorted(SUPPORTED_SUFFIXES)})"
        )
    root = Path(root) if root else path.parent
    try:
        source = str(path.relative_to(root)).replace("\\", "/")
    except ValueError:
        source = path.name
    return Document(
        doc_id=path.stem,
        source=source,
        fmt=suffix.lstrip("."),
        blocks=_LOADERS[suffix](path),
    )


def load_corpus(root: Path) -> list[Document]:
    """递归解析目录下所有受支持的文件,按路径排序以保证顺序可复现。"""
    root = Path(root)
    paths = sorted(
        p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in _LOADERS
    )
    return [load_document(p, root) for p in paths]
